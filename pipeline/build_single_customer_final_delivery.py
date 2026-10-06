# EN: Assemble verified local reports and catalogs; tests do not prove live capture completeness.
# 中文：组装已验证的本地报告和目录；测试不代表真实采集完整。
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
import xml.etree.ElementTree as ET
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SCHEMA = "okki.single_customer.combined_scope.v1"
DELIVERY_SCHEMA = "okki.single_customer.final_delivery_verification.v1"
FULL_EXTRACT_SCHEMA = "okki.single_customer_full_extract.v1"
WORKBOOK_PAYLOAD_SCHEMA = "okki.single_customer.delivery_workbook_payload.v1"
MANUAL_MANIFEST_SCHEMA = "okki.trade.manual_capture_manifest.v1"
MANUAL_RECEIPT_SCHEMA = "okki.trade.manual_capture_receipt.v1"
SHA256_RE = re.compile(r"^[A-Fa-f0-9]{64}$")
SESSION_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")
GENERATED_FILENAMES = {
    "combined_scope.v1.json",
    "customer_research_evidence.xlsx",
    "客户技术研究报告.docx",
    "客户技术研究报告.pdf",
    "workbook_verification.json",
    "delivery_verification.json",
    "checksums.sha256",
}


class DeliveryError(RuntimeError):
    pass


@dataclass(frozen=True)
class ArtifactRef:
    path: str
    bytes: int
    sha256: str | None

    def as_json(self) -> dict[str, Any]:
        value: dict[str, Any] = {"path": self.path, "bytes": self.bytes}
        if self.sha256 is not None:
            value["sha256"] = self.sha256
        return value


@dataclass
class SourceBundle:
    case_root: Path
    company_id: str
    processing_scope: dict[str, Any]
    processing_scope_path: Path
    full_extract: dict[str, Any]
    manual_trade: dict[str, Any]
    coverage_rows: list[list[Any]]
    relation_rows: list[list[Any]]
    evidence_rows: list[list[Any]]


def now_utc() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def canonical_json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode("utf-8")


def write_json_atomic(path: Path, value: Any) -> None:
    write_bytes_atomic(path, canonical_json_bytes(value))


def write_bytes_atomic(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    with temporary.open("wb") as handle:
        handle.write(payload)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary, path)


def load_json(path: Path, code: str) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise DeliveryError(code) from exc
    if not isinstance(value, dict):
        raise DeliveryError(code)
    return value


def require(condition: bool, code: str) -> None:
    if not condition:
        raise DeliveryError(code)


def resolve_within(root: Path, candidate: Path | str, code: str) -> Path:
    root = root.resolve()
    raw = Path(candidate)
    resolved = raw.resolve() if raw.is_absolute() else (root / raw).resolve()
    require(resolved == root or root in resolved.parents, code)
    return resolved


def relative_posix(path: Path, root: Path) -> str:
    return path.resolve().relative_to(root.resolve()).as_posix()


def artifact_ref(path: Path, root: Path, include_hash: bool = True) -> ArtifactRef:
    require(path.is_file(), "ARTIFACT_MISSING")
    return ArtifactRef(
        path=relative_posix(path, root),
        bytes=path.stat().st_size,
        sha256=sha256_file(path) if include_hash else None,
    )


def parse_nonnegative_int(value: Any, code: str) -> int:
    try:
        number = int(value)
    except (TypeError, ValueError) as exc:
        raise DeliveryError(code) from exc
    require(number >= 0, code)
    return number


def validate_sha(value: Any, code: str) -> str:
    text = str(value or "")
    require(SHA256_RE.fullmatch(text) is not None, code)
    return text.upper()


def validate_relative_artifact_path(value: Any, code: str) -> str:
    text = str(value or "").replace("\\", "/")
    pure = PurePosixPath(text)
    require(text != "" and not pure.is_absolute() and ".." not in pure.parts, code)
    return pure.as_posix()


def load_case_identity(case_root: Path, company_id: str) -> dict[str, Any]:
    require(SESSION_RE.fullmatch(company_id) is not None, "COMPANY_ID_INVALID")
    require(case_root.is_dir(), "CASE_ROOT_MISSING")
    identity = load_json(case_root / "case_identity.json", "CASE_IDENTITY_INVALID")
    require(str(identity.get("company_id") or "") == company_id, "CASE_IDENTITY_MISMATCH")
    require(
        case_root.name == f"company_{company_id}" or case_root.name.endswith(f"_{company_id}"),
        "CASE_DIRECTORY_IDENTITY_MISMATCH",
    )
    return identity


def load_processing_scope(case_root: Path, company_id: str, explicit: Path | None) -> tuple[dict[str, Any], Path]:
    path = resolve_within(
        case_root,
        explicit if explicit is not None else Path("manifests") / "processing_scope_latest.json",
        "PROCESSING_SCOPE_PATH_ESCAPE",
    )
    require(path.is_file(), "PROCESSING_SCOPE_MISSING")
    scope = load_json(path, "PROCESSING_SCOPE_INVALID")
    require(scope.get("schema") == 1, "PROCESSING_SCOPE_SCHEMA_INVALID")
    require(str(scope.get("capture_status") or "").casefold() == "pass", "PROCESSING_SCOPE_NOT_PASS")
    require(str(scope.get("company_id") or "") == company_id, "PROCESSING_SCOPE_COMPANY_MISMATCH")
    session_id = str(scope.get("default_session_id") or "")
    require(SESSION_RE.fullmatch(session_id) is not None, "PROCESSING_SCOPE_SESSION_INVALID")
    object_filter = scope.get("default_object_filter")
    require(
        isinstance(object_filter, dict)
        and object_filter.get("field") == "session_id"
        and str(object_filter.get("equals") or "") == session_id,
        "PROCESSING_SCOPE_FILTER_INVALID",
    )
    object_manifest_rel = validate_relative_artifact_path(scope.get("object_manifest"), "PROCESSING_SCOPE_OBJECT_MANIFEST_INVALID")
    object_manifest = resolve_within(case_root, object_manifest_rel, "PROCESSING_SCOPE_OBJECT_MANIFEST_ESCAPE")
    require(object_manifest.is_file(), "PROCESSING_SCOPE_OBJECT_MANIFEST_MISSING")
    terminal = resolve_within(
        case_root,
        Path("raw") / "sessions" / session_id / "final_status.json",
        "PROCESSING_SCOPE_TERMINAL_ESCAPE",
    )
    require(terminal.is_file(), "PROCESSING_SCOPE_TERMINAL_MISSING")
    terminal_value = load_json(terminal, "PROCESSING_SCOPE_TERMINAL_INVALID")
    errors = terminal_value.get("errors")
    metrics = terminal_value.get("metrics")
    require(
        terminal_value.get("phase") == "complete"
        and isinstance(errors, list)
        and not errors
        and isinstance(metrics, dict)
        and metrics.get("reconciliation_failures") == 0,
        "PROCESSING_SCOPE_TERMINAL_NOT_PASS",
    )
    return scope, path


def read_csv_rows(path: Path, required_headers: set[str], code: str) -> list[dict[str, str]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            require(reader.fieldnames is not None and required_headers.issubset(set(reader.fieldnames)), code)
            return [dict(row) for row in reader]
    except (OSError, UnicodeError, csv.Error) as exc:
        raise DeliveryError(code) from exc


def load_full_extract(case_root: Path, company_id: str) -> dict[str, Any]:
    root = resolve_within(case_root, Path("derived") / "full_extract", "FULL_EXTRACT_PATH_ESCAPE")
    require(root.is_dir(), "FULL_EXTRACT_MISSING")
    verification_path = root / "verification.json"
    verification = load_json(verification_path, "FULL_EXTRACT_VERIFICATION_INVALID")
    require(verification.get("schema") == FULL_EXTRACT_SCHEMA, "FULL_EXTRACT_SCHEMA_INVALID")
    require(str(verification.get("company_id") or "") == company_id, "FULL_EXTRACT_COMPANY_MISMATCH")
    require(verification.get("status") == "PASS", "FULL_EXTRACT_NOT_PASS")

    paths = {
        "file_inventory": root / "file_inventory.csv",
        "evidence_index": root / "evidence_index.csv",
        "media_inventory": root / "media_inventory.csv",
        "coverage": root / "coverage.csv",
        "relationships": root / "relationship_graph.jsonl",
        "full_text": root / "full_text.sqlite3",
        "verification": verification_path,
    }
    for name, path in paths.items():
        require(path.is_file(), f"FULL_EXTRACT_{name.upper()}_MISSING")

    files = read_csv_rows(
        paths["file_inventory"],
        {"source_rel", "sha256", "bytes", "magic", "category", "mime", "active", "prepared_at"},
        "FULL_EXTRACT_FILE_INVENTORY_INVALID",
    )
    tasks = read_csv_rows(
        paths["evidence_index"],
        {"task_id", "source_rel", "source_sha256", "task_kind", "state", "output_rel", "result_sha256", "attempts", "error_code", "updated_at"},
        "FULL_EXTRACT_EVIDENCE_INDEX_INVALID",
    )
    media = read_csv_rows(
        paths["media_inventory"],
        {"source_rel", "sha256", "bytes", "category", "ffprobe", "dual_asr_diarization", "video_frame_ocr"},
        "FULL_EXTRACT_MEDIA_INVENTORY_INVALID",
    )
    coverage = read_csv_rows(
        paths["coverage"],
        {"scope", "name", "total", "completed", "pending", "failed", "bytes"},
        "FULL_EXTRACT_COVERAGE_INVALID",
    )

    relation_rows: list[list[Any]] = []
    relation_nodes = 0
    relation_edges = 0
    try:
        with paths["relationships"].open("r", encoding="utf-8") as handle:
            for line in handle:
                if not line.strip():
                    continue
                value = json.loads(line)
                require(isinstance(value, dict), "FULL_EXTRACT_RELATION_RECORD_INVALID")
                kind = value.get("record_type")
                if kind == "node":
                    relation_nodes += 1
                    relation_rows.append([
                        "node",
                        str(value.get("node_id") or ""),
                        str(value.get("entity_type") or ""),
                        "",
                        "",
                        "",
                        "",
                        "",
                        "原始 value 已裁剪，未进入交付",
                    ])
                elif kind == "edge":
                    relation_edges += 1
                    relation_rows.append([
                        "edge",
                        str(value.get("edge_id") or ""),
                        "",
                        str(value.get("from") or ""),
                        str(value.get("to") or ""),
                        str(value.get("relation") or ""),
                        str(value.get("source_rel") or ""),
                        str(value.get("json_pointer") or ""),
                        "端点为哈希 node_id；未导出节点原值",
                    ])
                else:
                    raise DeliveryError("FULL_EXTRACT_RELATION_RECORD_TYPE_INVALID")
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise DeliveryError("FULL_EXTRACT_RELATION_GRAPH_INVALID") from exc

    require(len(files) == parse_nonnegative_int(verification.get("evidence_files"), "FULL_EXTRACT_EVIDENCE_FILE_COUNT_INVALID"), "FULL_EXTRACT_EVIDENCE_FILE_COUNT_MISMATCH")
    require(len(tasks) == parse_nonnegative_int(verification.get("tasks"), "FULL_EXTRACT_TASK_COUNT_INVALID"), "FULL_EXTRACT_TASK_COUNT_MISMATCH")
    require(relation_nodes + relation_edges == parse_nonnegative_int(verification.get("relation_rows"), "FULL_EXTRACT_RELATION_COUNT_INVALID"), "FULL_EXTRACT_RELATION_COUNT_MISMATCH")
    require(parse_nonnegative_int(verification.get("pending"), "FULL_EXTRACT_PENDING_INVALID") == 0, "FULL_EXTRACT_PENDING_NONZERO")
    require(parse_nonnegative_int(verification.get("failed"), "FULL_EXTRACT_FAILED_INVALID") == 0, "FULL_EXTRACT_FAILED_NONZERO")

    evidence_rows: list[list[Any]] = []
    for row in files:
        evidence_rows.append([
            "full_extract_source_file",
            "ACTIVE" if str(row.get("active")) == "1" else "INACTIVE",
            row.get("source_rel", ""),
            validate_sha(row.get("sha256"), "FULL_EXTRACT_FILE_SHA_INVALID"),
            parse_nonnegative_int(row.get("bytes"), "FULL_EXTRACT_FILE_BYTES_INVALID"),
            row.get("category", ""),
            "",
            "",
            f"magic={row.get('magic', '')}; mime={row.get('mime', '')}",
        ])
    for row in tasks:
        source_sha = validate_sha(row.get("source_sha256"), "FULL_EXTRACT_TASK_SOURCE_SHA_INVALID")
        result_sha = str(row.get("result_sha256") or "")
        if result_sha:
            result_sha = validate_sha(result_sha, "FULL_EXTRACT_TASK_RESULT_SHA_INVALID")
        evidence_rows.append([
            "full_extract_task",
            str(row.get("state") or ""),
            row.get("source_rel", ""),
            source_sha,
            "",
            row.get("task_kind", ""),
            row.get("output_rel", ""),
            result_sha,
            f"attempts={row.get('attempts', '')}; error={row.get('error_code', '')}",
        ])
    for row in media:
        evidence_rows.append([
            "full_extract_media_status",
            "INDEXED",
            row.get("source_rel", ""),
            validate_sha(row.get("sha256"), "FULL_EXTRACT_MEDIA_SHA_INVALID"),
            parse_nonnegative_int(row.get("bytes"), "FULL_EXTRACT_MEDIA_BYTES_INVALID"),
            row.get("category", ""),
            "",
            "",
            f"ffprobe={row.get('ffprobe', '')}; dual_asr={row.get('dual_asr_diarization', '')}; frame_ocr={row.get('video_frame_ocr', '')}",
        ])

    coverage_rows: list[list[Any]] = []
    for row in coverage:
        total = parse_nonnegative_int(row.get("total"), "FULL_EXTRACT_COVERAGE_TOTAL_INVALID")
        completed = parse_nonnegative_int(row.get("completed"), "FULL_EXTRACT_COVERAGE_COMPLETED_INVALID")
        pending = parse_nonnegative_int(row.get("pending"), "FULL_EXTRACT_COVERAGE_PENDING_INVALID")
        failed = parse_nonnegative_int(row.get("failed"), "FULL_EXTRACT_COVERAGE_FAILED_INVALID")
        status = "PASS" if total == completed and pending == 0 and failed == 0 else "INCOMPLETE"
        coverage_rows.append([
            "FULL_EXTRACT",
            row.get("scope", ""),
            row.get("name", ""),
            total,
            completed,
            pending,
            failed,
            status,
            relative_posix(paths["coverage"], case_root),
            f"bytes={row.get('bytes', '')}",
        ])

    refs = {name: artifact_ref(path, case_root, include_hash=name != "full_text").as_json() for name, path in paths.items()}
    refs["full_text"]["semantic_read"] = False
    refs["full_text"]["note"] = "仅确认文件存在与字节数，最终交付生成器未打开全文索引"
    return {
        "root": root,
        "verification_value": verification,
        "refs": refs,
        "files": files,
        "tasks": tasks,
        "media": media,
        "coverage_rows": coverage_rows,
        "relation_rows": relation_rows,
        "evidence_rows": evidence_rows,
        "counts": {
            "raw_files": len(files),
            "tasks": len(tasks),
            "relation_nodes": relation_nodes,
            "relation_edges": relation_edges,
            "fts_rows": parse_nonnegative_int(verification.get("fts_rows"), "FULL_EXTRACT_FTS_ROWS_INVALID"),
        },
    }


def discover_manual_receipt(case_root: Path) -> Path:
    root = resolve_within(case_root, Path("evidence") / "manual_trade", "MANUAL_TRADE_ROOT_ESCAPE")
    require(root.is_dir(), "MANUAL_TRADE_ROOT_MISSING")
    candidates = sorted(root.glob("*/trade_manual_capture_receipt.json"), key=lambda p: (p.stat().st_mtime_ns, p.as_posix()))
    require(bool(candidates), "MANUAL_TRADE_RECEIPT_MISSING")
    return candidates[-1]


def load_manual_trade(
    case_root: Path,
    company_id: str,
    explicit_receipt: Path | None,
    explicit_manifest: Path | None,
    verify_artifact_hashes: bool,
) -> dict[str, Any]:
    receipt_path = resolve_within(
        case_root,
        explicit_receipt if explicit_receipt is not None else discover_manual_receipt(case_root),
        "MANUAL_TRADE_RECEIPT_ESCAPE",
    )
    require(receipt_path.is_file(), "MANUAL_TRADE_RECEIPT_MISSING")
    receipt = load_json(receipt_path, "MANUAL_TRADE_RECEIPT_INVALID")
    require(receipt.get("schema") == MANUAL_RECEIPT_SCHEMA, "MANUAL_TRADE_RECEIPT_SCHEMA_INVALID")
    require(str(receipt.get("company_id") or "") == company_id, "MANUAL_TRADE_RECEIPT_COMPANY_MISMATCH")
    receipt_status = str(receipt.get("status") or "")
    require(receipt_status.startswith("PASS"), "MANUAL_TRADE_RECEIPT_NOT_PASS")
    session = str(receipt.get("session") or "")
    require(SESSION_RE.fullmatch(session) is not None, "MANUAL_TRADE_SESSION_INVALID")

    manifest_candidate: Path | str
    if explicit_manifest is not None:
        manifest_candidate = explicit_manifest
    else:
        manifest_candidate = str(receipt.get("manifest_path") or "")
    manifest_path = resolve_within(case_root, manifest_candidate, "MANUAL_TRADE_MANIFEST_ESCAPE")
    require(manifest_path.is_file(), "MANUAL_TRADE_MANIFEST_MISSING")
    manifest_bytes = parse_nonnegative_int(receipt.get("manifest_bytes"), "MANUAL_TRADE_MANIFEST_BYTES_INVALID")
    manifest_sha = validate_sha(receipt.get("manifest_sha256"), "MANUAL_TRADE_MANIFEST_SHA_INVALID")
    require(manifest_path.stat().st_size == manifest_bytes, "MANUAL_TRADE_MANIFEST_BYTES_MISMATCH")
    require(sha256_file(manifest_path) == manifest_sha, "MANUAL_TRADE_MANIFEST_SHA_MISMATCH")

    manifest = load_json(manifest_path, "MANUAL_TRADE_MANIFEST_INVALID")
    require(manifest.get("schema") == MANUAL_MANIFEST_SCHEMA, "MANUAL_TRADE_MANIFEST_SCHEMA_INVALID")
    require(str(manifest.get("company_id") or "") == company_id, "MANUAL_TRADE_MANIFEST_COMPANY_MISMATCH")
    require(str(manifest.get("session") or "") == session, "MANUAL_TRADE_MANIFEST_SESSION_MISMATCH")
    require(str(manifest.get("status") or "") == receipt_status, "MANUAL_TRADE_STATUS_MISMATCH")
    require(manifest.get("counts") == receipt.get("counts"), "MANUAL_TRADE_COUNTS_MISMATCH")
    counts = manifest.get("counts")
    require(isinstance(counts, dict), "MANUAL_TRADE_COUNTS_INVALID")
    artifacts = manifest.get("artifacts")
    require(isinstance(artifacts, list), "MANUAL_TRADE_ARTIFACTS_INVALID")

    evidence_rows: list[list[Any]] = []
    seen: set[str] = set()
    total_bytes = 0
    rehashed = 0
    for item in artifacts:
        require(isinstance(item, dict), "MANUAL_TRADE_ARTIFACT_INVALID")
        rel = validate_relative_artifact_path(item.get("relative_path"), "MANUAL_TRADE_ARTIFACT_PATH_INVALID")
        require(rel not in seen, "MANUAL_TRADE_ARTIFACT_DUPLICATE")
        seen.add(rel)
        claimed_bytes = parse_nonnegative_int(item.get("bytes"), "MANUAL_TRADE_ARTIFACT_BYTES_INVALID")
        claimed_sha = validate_sha(item.get("sha256"), "MANUAL_TRADE_ARTIFACT_SHA_INVALID")
        artifact_path = resolve_within(manifest_path.parent, rel, "MANUAL_TRADE_ARTIFACT_ESCAPE")
        require(artifact_path.is_file(), "MANUAL_TRADE_ARTIFACT_MISSING")
        require(artifact_path.stat().st_size == claimed_bytes, "MANUAL_TRADE_ARTIFACT_BYTES_MISMATCH")
        if verify_artifact_hashes:
            require(sha256_file(artifact_path) == claimed_sha, "MANUAL_TRADE_ARTIFACT_SHA_MISMATCH")
            rehashed += 1
        total_bytes += claimed_bytes
        evidence_rows.append([
            "manual_trade_artifact",
            "HASH_VERIFIED" if verify_artifact_hashes else "RECEIPT_BOUND",
            relative_posix(artifact_path, case_root),
            claimed_sha,
            claimed_bytes,
            "manual_trade",
            "",
            "",
            "未读取正文；仅验证路径/字节" + ("/SHA256" if verify_artifact_hashes else "；SHA256由manifest+receipt绑定"),
        ])

    require(len(artifacts) == parse_nonnegative_int(counts.get("artifact_files"), "MANUAL_TRADE_ARTIFACT_COUNT_INVALID"), "MANUAL_TRADE_ARTIFACT_COUNT_MISMATCH")
    require(total_bytes == parse_nonnegative_int(counts.get("artifact_bytes"), "MANUAL_TRADE_ARTIFACT_BYTES_TOTAL_INVALID"), "MANUAL_TRADE_ARTIFACT_BYTES_TOTAL_MISMATCH")
    source_gaps = manifest.get("source_gaps") or []
    require(isinstance(source_gaps, list), "MANUAL_TRADE_SOURCE_GAPS_INVALID")
    gap_rows = []
    for gap in source_gaps:
        require(isinstance(gap, dict), "MANUAL_TRADE_SOURCE_GAP_INVALID")
        gap_rows.append({
            "candidate_index": parse_nonnegative_int(gap.get("candidate_index"), "MANUAL_TRADE_SOURCE_GAP_CANDIDATE_INVALID"),
            "code": str(gap.get("code") or "UNSPECIFIED_GAP"),
            "observed_numeric_token_count": len(gap.get("observed_numeric_tokens") or []) if isinstance(gap.get("observed_numeric_tokens") or [], list) else 0,
        })

    coverage_rows: list[list[Any]] = []
    for key in sorted(counts):
        number = parse_nonnegative_int(counts.get(key), "MANUAL_TRADE_COUNT_INVALID")
        coverage_rows.append([
            "MANUAL_TRADE",
            "receipt_count",
            key,
            number,
            number,
            0,
            0,
            receipt_status,
            relative_posix(receipt_path, case_root),
            "manifest 与 receipt 计数一致",
        ])
    for gap in gap_rows:
        coverage_rows.append([
            "MANUAL_TRADE",
            "source_gap",
            gap["code"],
            1,
            0,
            1,
            0,
            "SOURCE_GAP",
            relative_posix(manifest_path, case_root),
            f"candidate_index={gap['candidate_index']}; numeric_token_count={gap['observed_numeric_token_count']}",
        ])

    return {
        "manifest": manifest,
        "receipt": receipt,
        "manifest_path": manifest_path,
        "receipt_path": receipt_path,
        "status": receipt_status,
        "session": session,
        "counts": {key: parse_nonnegative_int(value, "MANUAL_TRADE_COUNT_INVALID") for key, value in counts.items()},
        "source_gaps": gap_rows,
        "evidence_rows": evidence_rows,
        "coverage_rows": coverage_rows,
        "validation": {
            "artifact_paths_and_sizes_verified": len(artifacts),
            "artifact_sha256_rehashed": rehashed,
            "mode": "FULL_HASH" if verify_artifact_hashes else "RECEIPT_BOUND_METADATA_ONLY",
        },
    }


def build_source_bundle(
    case_root: Path,
    company_id: str,
    processing_scope_path: Path | None,
    trade_receipt_path: Path | None,
    trade_manifest_path: Path | None,
    verify_manual_hashes: bool,
) -> SourceBundle:
    case_root = case_root.resolve()
    load_case_identity(case_root, company_id)
    scope, scope_path = load_processing_scope(case_root, company_id, processing_scope_path)
    full_extract = load_full_extract(case_root, company_id)
    manual = load_manual_trade(case_root, company_id, trade_receipt_path, trade_manifest_path, verify_manual_hashes)

    scope_ref = artifact_ref(scope_path, case_root).as_json()
    coverage_rows = [[
        "AUTO7",
        "processing_scope",
        "通过完整性门的默认会话",
        1,
        1,
        0,
        0,
        "PASS",
        scope_ref["path"],
        f"session_id={scope.get('default_session_id')}",
    ], *full_extract["coverage_rows"], *manual["coverage_rows"]]

    evidence_rows = [
        ["AUTO7_processing_scope", "PASS", scope_ref["path"], scope_ref["sha256"], scope_ref["bytes"], "processing_scope", "", "", "自动采集范围绑定"],
        [
            "manual_trade_manifest",
            manual["status"],
            relative_posix(manual["manifest_path"], case_root),
            sha256_file(manual["manifest_path"]),
            manual["manifest_path"].stat().st_size,
            "manual_trade_manifest",
            "",
            "",
            "手工贸易证据清单",
        ],
        [
            "manual_trade_receipt",
            manual["status"],
            relative_posix(manual["receipt_path"], case_root),
            sha256_file(manual["receipt_path"]),
            manual["receipt_path"].stat().st_size,
            "manual_trade_receipt",
            "",
            "",
            "手工贸易采集回执",
        ],
        *full_extract["evidence_rows"],
        *manual["evidence_rows"],
    ]
    return SourceBundle(
        case_root=case_root,
        company_id=company_id,
        processing_scope=scope,
        processing_scope_path=scope_path,
        full_extract=full_extract,
        manual_trade=manual,
        coverage_rows=coverage_rows,
        relation_rows=full_extract["relation_rows"],
        evidence_rows=evidence_rows,
    )


def combined_scope(bundle: SourceBundle, generated_at: str) -> dict[str, Any]:
    manual_status = bundle.manual_trade["status"]
    scope_status = "PASS_WITH_SOURCE_GAPS" if manual_status == "PASS_WITH_SOURCE_GAPS" or bundle.manual_trade["source_gaps"] else "PASS"
    return {
        "schema": SCHEMA,
        "generated_at_utc": generated_at,
        "company_id": bundle.company_id,
        "status": scope_status,
        "automatic_capture": {
            "channel": "AUTO7",
            "status": "PASS",
            "session_id": bundle.processing_scope.get("default_session_id"),
            "processing_scope": artifact_ref(bundle.processing_scope_path, bundle.case_root).as_json(),
            "object_manifest": bundle.processing_scope.get("object_manifest"),
            "default_object_filter": bundle.processing_scope.get("default_object_filter"),
        },
        "full_extract": {
            "status": "PASS",
            "verification": bundle.full_extract["refs"]["verification"],
            "indexes": {
                key: bundle.full_extract["refs"][key]
                for key in ("file_inventory", "evidence_index", "media_inventory", "coverage", "relationships", "full_text")
            },
            "counts": bundle.full_extract["counts"],
        },
        "manual_trade": {
            "status": manual_status,
            "session": bundle.manual_trade["session"],
            "manifest": artifact_ref(bundle.manual_trade["manifest_path"], bundle.case_root).as_json(),
            "receipt": artifact_ref(bundle.manual_trade["receipt_path"], bundle.case_root).as_json(),
            "counts": bundle.manual_trade["counts"],
            "source_gaps": bundle.manual_trade["source_gaps"],
            "artifact_validation": bundle.manual_trade["validation"],
        },
        "privacy_boundary": {
            "customer_body_read_by_delivery_generator": False,
            "full_text_fts_opened": False,
            "relationship_node_values_exported": False,
            "relationship_endpoints": "hashed_node_ids_only",
            "console_output": "status_counts_and_paths_only",
        },
    }


def workbook_payload(bundle: SourceBundle, scope: dict[str, Any]) -> dict[str, Any]:
    return {
        "schema": WORKBOOK_PAYLOAD_SCHEMA,
        "company_id": bundle.company_id,
        "scope_status": scope["status"],
        "auto7_session_id": scope["automatic_capture"]["session_id"],
        "full_extract_status": scope["full_extract"]["status"],
        "manual_trade_session": scope["manual_trade"]["session"],
        "manual_trade_status": scope["manual_trade"]["status"],
        "counts": {
            **bundle.full_extract["counts"],
            "manual_trade_artifacts": bundle.manual_trade["counts"].get("artifact_files", 0),
            "manual_trade_artifact_bytes": bundle.manual_trade["counts"].get("artifact_bytes", 0),
        },
        "coverage_rows": bundle.coverage_rows,
        "relation_rows": bundle.relation_rows,
        "evidence_rows": bundle.evidence_rows,
    }


def locate_bundled_runtime(explicit_node: Path | None, explicit_modules: Path | None) -> tuple[Path, Path]:
    # EN: Retain the API name for compatibility; discover standard Node and repository npm dependencies.
    # 中文：保留接口名以兼容调用方，但使用标准 Node 与仓库 npm 依赖。
    selected = str(explicit_node) if explicit_node else os.environ.get("CRM_NODE") or shutil.which("node")
    require(bool(selected), "NODE_EXECUTABLE_MISSING")
    node = Path(str(selected)).resolve()
    modules = (explicit_modules or Path(__file__).resolve().parents[1] / "node_modules").resolve()
    require(node.is_file(), "NODE_EXECUTABLE_MISSING")
    require((modules / "exceljs").is_dir(), "EXCELJS_DEPENDENCY_MISSING_RUN_NPM_CI")
    return node, modules


def patch_xlsx_freeze_panes(path: Path) -> int:
    """Normalize view-only freeze metadata in the exported workbook.

    This patch never reads or changes cell values, formulas, tables, styles, or
    shared strings.  It only inserts a frozen pane definition into each
    worksheet view after the open-source builder has authored the workbook.
    中文：仅规范冻结窗格，不读取或修改数据、公式、样式或表格。
    """
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", spreadsheet_ns)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.freeze.tmp")
    patched = 0
    try:
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w") as target_zip:
            for info in source.infolist():
                payload = source.read(info.filename)
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", info.filename):
                    root = ET.fromstring(payload)
                    views = root.find(f"{{{spreadsheet_ns}}}sheetViews")
                    view = views.find(f"{{{spreadsheet_ns}}}sheetView") if views is not None else None
                    require(view is not None, "WORKBOOK_SHEET_VIEW_MISSING")
                    for pane in list(view.findall(f"{{{spreadsheet_ns}}}pane")):
                        view.remove(pane)
                    pane = ET.Element(
                        f"{{{spreadsheet_ns}}}pane",
                        {
                            "xSplit": "1",
                            "ySplit": "4",
                            "topLeftCell": "B5",
                            "activePane": "bottomRight",
                            "state": "frozen",
                        },
                    )
                    view.insert(0, pane)
                    payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                    patched += 1
                target_zip.writestr(info, payload)
        require(patched > 0, "WORKBOOK_FREEZE_PATCH_NO_SHEETS")
        os.replace(temporary, path)
        return patched
    except (OSError, zipfile.BadZipFile, ET.ParseError) as exc:
        raise DeliveryError("WORKBOOK_FREEZE_PATCH_FAILED") from exc
    finally:
        if temporary.exists():
            temporary.unlink()


def build_workbook(
    payload: dict[str, Any],
    output_path: Path,
    preview_dir: Path,
    verification_path: Path,
    node_exe: Path | None,
    node_modules: Path | None,
) -> None:
    node, _modules = locate_bundled_runtime(node_exe, node_modules)
    builder_source = Path(__file__).with_name("build_single_customer_delivery_workbook.mjs")
    require(builder_source.is_file(), "WORKBOOK_BUILDER_MISSING")
    with tempfile.TemporaryDirectory(prefix="okki_final_delivery_workbook_") as temp_name:
        temp = Path(temp_name)
        # EN: Run the original repository module so relative libraries remain resolvable.
        # 中文：直接运行仓库模块，确保相对库路径可解析；不复制脚本或创建依赖链接。
        payload_path = temp / "payload.json"
        payload_path.write_bytes(canonical_json_bytes(payload))
        completed = subprocess.run(
            [
                str(node), str(builder_source),
                "--input", str(payload_path),
                "--output", str(output_path),
                "--preview-dir", str(preview_dir),
                "--verification", str(verification_path),
            ],
            cwd=builder_source.parent.parent,
            stdin=subprocess.DEVNULL,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
        )
        if completed.returncode != 0:
            code = next((line.strip() for line in completed.stderr.splitlines() if re.fullmatch(r"[A-Z0-9_]+", line.strip())), "WORKBOOK_BUILD_FAILED")
            raise DeliveryError(code)
    require(output_path.is_file() and output_path.stat().st_size > 0, "WORKBOOK_OUTPUT_MISSING")
    frozen_sheets = patch_xlsx_freeze_panes(output_path)
    verification = load_json(verification_path, "WORKBOOK_VERIFICATION_INVALID")
    require(verification.get("status") == "PASS", "WORKBOOK_VERIFICATION_NOT_PASS")
    require(verification.get("privacy", {}).get("relation_node_values_exported") is False, "WORKBOOK_PRIVACY_INVALID")
    verification["output"] = artifact_ref(output_path, output_path.parent).as_json()
    verification["post_export_view_fix"] = {
        "reason": "normalize frozen pane XML after open-source export",
        "scope": "worksheet view metadata only; cell values/formulas/styles/tables untouched",
        "frozen_sheets": frozen_sheets,
        "top_left_cell": "B5",
    }
    write_json_atomic(verification_path, verification)


def set_run_font(run: Any, size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell: Any, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell: Any, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    margins = tcpr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tcpr.append(margins)
    for side, width in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table: Any, widths: list[int], indent: int = 120) -> None:
    require(sum(widths) == 9360, "DOCX_TABLE_WIDTH_INVALID")
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.first_child_found_in("w:tblW")
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), "9360")
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.first_child_found_in("w:tblInd")
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(indent))
    tblind.set(qn("w:type"), "dxa")
    layout = tblpr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.first_child_found_in("w:tcW")
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths[index]))
            tcw.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_header_row(row: Any) -> None:
    trpr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    trpr.append(marker)


def style_table(table: Any, widths: list[int]) -> None:
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    repeat_header_row(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            if row_index == 0:
                set_cell_shading(cell, "E8EEF5")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(run, size=9.5, color="0B2545" if row_index == 0 else "1F2937", bold=row_index == 0)


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[Any]], widths: list[int]) -> Any:
    row_values = [list(row) for row in rows]
    if not row_values:
        row_values = [["-" for _ in headers]]
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in row_values:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value if value is not None else "")
    style_table(table, widths)
    return table


def add_heading(doc: Document, text: str, level: int) -> Any:
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    return paragraph


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> Any:
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        set_run_font(lead, bold=True)
        body = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(body)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, size=8.5, color="6B7280")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("OKKI | 单客户完整研究交付"), size=8.5, color="6B7280")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run("内部技术研究交付 | 第 "), size=8.5, color="6B7280")
    add_page_field(footer)
    set_run_font(footer.add_run(" 页"), size=8.5, color="6B7280")


def build_docx(bundle: SourceBundle, scope: dict[str, Any], output_path: Path) -> None:
    doc = Document()
    configure_document(doc)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "单客户完整研究技术报告"
    doc.core_properties.subject = "仅包含证据索引、哈希、计数和隐私裁剪关系结构"
    doc.core_properties.keywords = "OKKI, evidence, metadata-only, privacy"

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run("单客户完整研究技术报告"), size=23, color="0B2545", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(subtitle.add_run("证据范围、覆盖矩阵与关系结构（隐私裁剪版）"), size=13, color="4B5563")
    add_table(
        doc,
        ["元数据项", "值"],
        [
            ["公司 ID", bundle.company_id],
            ["联合范围状态", scope["status"]],
            ["AUTO7 会话", scope["automatic_capture"]["session_id"]],
            ["手工贸易会话", scope["manual_trade"]["session"]],
            ["生成时间（UTC）", scope["generated_at_utc"]],
        ],
        [2700, 6660],
    )

    add_heading(doc, "1. 结论与交付边界", 1)
    add_body(doc, f"结论：自动采集范围与完整提取均通过完整性门；手工贸易证据状态为 {scope['manual_trade']['status']}；联合范围状态为 {scope['status']}。", "结论：")
    add_body(doc, "隐私边界：本生成器没有打开全文 FTS，没有提取或输出客户正文，也没有把关系节点的原始 value 写入 DOCX 或 Excel。关系端点仅保留哈希 node_id。", "隐私边界：")
    add_body(doc, "Excel 安全：所有写入单元格的外部文本均对 =、+、-、@ 及前导空白危险前缀做文本转义。", "Excel 安全：")

    add_heading(doc, "2. 输入证据绑定", 1)
    input_rows = [
        ["AUTO7 processing_scope", scope["automatic_capture"]["processing_scope"]["path"], scope["automatic_capture"]["processing_scope"]["bytes"], scope["automatic_capture"]["processing_scope"]["sha256"]],
        ["完整提取 verification", scope["full_extract"]["verification"]["path"], scope["full_extract"]["verification"]["bytes"], scope["full_extract"]["verification"]["sha256"]],
        ["手工贸易 manifest", scope["manual_trade"]["manifest"]["path"], scope["manual_trade"]["manifest"]["bytes"], scope["manual_trade"]["manifest"]["sha256"]],
        ["手工贸易 receipt", scope["manual_trade"]["receipt"]["path"], scope["manual_trade"]["receipt"]["bytes"], scope["manual_trade"]["receipt"]["sha256"]],
    ]
    add_table(doc, ["来源", "相对路径", "字节", "SHA-256"], input_rows, [1800, 3420, 1140, 3000])

    add_heading(doc, "3. 覆盖与完成度", 1)
    counts = bundle.full_extract["counts"]
    manual_counts = bundle.manual_trade["counts"]
    add_table(
        doc,
        ["指标", "数量", "状态"],
        [
            ["完整提取原始文件", counts["raw_files"], "PASS"],
            ["完整提取任务", counts["tasks"], "PASS"],
            ["全文索引文档行", counts["fts_rows"], "仅计数，未打开正文"],
            ["关系节点", counts["relation_nodes"], "value 已裁剪"],
            ["关系边", counts["relation_edges"], "哈希端点"],
            ["手工贸易证据文件", manual_counts.get("artifact_files", 0), bundle.manual_trade["status"]],
            ["手工贸易证据字节", manual_counts.get("artifact_bytes", 0), bundle.manual_trade["validation"]["mode"]],
        ],
        [3900, 1740, 3720],
    )
    add_body(doc, "完整逐项覆盖记录、待处理/失败计数和证据引用位于同目录 Excel 的“覆盖矩阵”工作表。")

    add_heading(doc, "4. 关系结构摘要", 1)
    node_types = Counter(row[2] for row in bundle.relation_rows if row[0] == "node" and row[2])
    edge_types = Counter(row[5] for row in bundle.relation_rows if row[0] == "edge" and row[5])
    relation_summary = [["节点类型", key, value] for key, value in sorted(node_types.items())]
    relation_summary.extend([["关系类型", key, value] for key, value in sorted(edge_types.items())])
    add_table(doc, ["类别", "名称", "数量"], relation_summary, [2100, 5160, 2100])
    add_body(doc, "关系台账仅保留 record_id、entity_type、from/to 哈希端点、relation、相对来源路径和 JSON Pointer。节点 value 永不进入交付。")

    add_heading(doc, "5. 手工贸易证据与已知缺口", 1)
    gap_rows = [
        [gap["candidate_index"], gap["code"], gap["observed_numeric_token_count"]]
        for gap in bundle.manual_trade["source_gaps"]
    ]
    if gap_rows:
        add_table(doc, ["候选序号", "缺口代码", "已观察数值令牌数"], gap_rows, [1800, 4860, 2700])
        add_body(doc, "说明：PASS_WITH_SOURCE_GAPS 表示证据采集和落盘本身通过，但上游页面存在已明确记录的来源缺口；该缺口不能被解释为无贸易记录。", "说明：")
    else:
        add_body(doc, "未记录手工贸易来源缺口。")

    add_heading(doc, "6. 方法、复核与限制", 1)
    add_body(doc, "范围绑定：AUTO7 processing_scope 固定默认通过会话；完整提取依赖 verification.json 的 PASS 状态；手工贸易 manifest 必须由 receipt 的字节数和 SHA-256 精确绑定。", "范围绑定：")
    add_body(doc, "完整性：交付生成阶段校验索引 schema、公司 ID、计数、路径边界、manifest/receipt 一致性和手工贸易证据文件存在性/字节数。可用 --verify-manual-artifact-hashes 追加逐文件 SHA-256 复核。", "完整性：")
    add_body(doc, "限制：本报告是技术研究交付索引，不是客户事实摘要；它不会把模糊匹配贸易候选自动升级为已确认客户事实，也不会从正文中生成销售判断。", "限制：")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(f".{output_path.name}.{uuid.uuid4().hex}.tmp")
    doc.save(temporary)
    os.replace(temporary, output_path)
    verify_docx_structure(output_path)


def verify_docx_structure(path: Path) -> None:
    require(path.is_file() and path.stat().st_size > 0, "DOCX_OUTPUT_MISSING")
    try:
        with zipfile.ZipFile(path) as archive:
            require(archive.testzip() is None, "DOCX_ZIP_CORRUPT")
            names = set(archive.namelist())
            require("word/document.xml" in names and "word/styles.xml" in names, "DOCX_STRUCTURE_INCOMPLETE")
    except (OSError, zipfile.BadZipFile) as exc:
        raise DeliveryError("DOCX_STRUCTURE_INVALID") from exc


def export_pdf_with_word(docx_path: Path, output_dir: Path, company_id: str, timeout_seconds: int) -> tuple[Path, Path]:
    script = Path(__file__).with_name("export_docx_pdf.ps1").resolve()
    worker = Path(__file__).with_name("word_com_export_worker.ps1").resolve()
    powershell = Path(os.environ.get("SystemRoot", r"C:\Windows")) / "System32" / "WindowsPowerShell" / "v1.0" / "powershell.exe"
    require(script.is_file() and worker.is_file() and powershell.is_file(), "WORD_EXPORT_RUNTIME_MISSING")
    pdf = output_dir / "客户技术研究报告.pdf"
    word_dir = output_dir / "qa" / "word"
    status = word_dir / "word_pdf_export_receipt.json"
    word_dir.mkdir(parents=True, exist_ok=True)
    ascii_input = word_dir / "report_input.docx"
    ascii_output = word_dir / "report_output.pdf"
    require(not ascii_input.exists() and not ascii_output.exists() and not pdf.exists(), "WORD_EXPORT_OUTPUT_PREEXISTS")
    shutil.copy2(docx_path, ascii_input)
    require(sha256_file(ascii_input) == sha256_file(docx_path), "WORD_EXPORT_ASCII_STAGE_HASH_MISMATCH")
    run_id = "final_delivery_" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    command = [
        str(powershell), "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(script),
        "-CustomerKey", company_id,
        "-RunId", run_id,
        "-InputDocx", str(ascii_input),
        "-OutputPdf", str(ascii_output),
        "-StatusJson", str(status),
        "-ExpectedInputSha256", sha256_file(ascii_input),
        "-ExpectedParentSha256", sha256_file(script),
        "-ExpectedWorkerSha256", sha256_file(worker),
        "-ExpectedPowerShellSha256", sha256_file(powershell),
        "-WorkerTimeoutSec", str(timeout_seconds),
    ]
    child_env = os.environ.copy()
    # EN: A PowerShell 7 host can prepend incompatible modules to PSModulePath. A
    # directly spawned Windows PowerShell 5.1 then finds the incompatible v7
    # Microsoft.PowerShell.Utility first and cannot auto-load Get-FileHash.
    # Give this isolated Office path only Windows PowerShell module roots.
    # 中文：Office 隔离子进程只使用 Windows PowerShell 的模块目录。
    program_files = Path(os.environ.get("ProgramFiles", r"C:\Program Files"))
    child_env["PSModulePath"] = os.pathsep.join(
        [
            str(Path.home() / "Documents" / "WindowsPowerShell" / "Modules"),
            str(program_files / "WindowsPowerShell" / "Modules"),
            str(Path(os.environ.get("SystemRoot", r"C:\Windows")) / "System32" / "WindowsPowerShell" / "v1.0" / "Modules"),
        ]
    )
    completed = subprocess.run(
        command,
        env=child_env,
        stdin=subprocess.DEVNULL,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
    )
    if completed.returncode != 0:
        receipt = load_json(status, "WORD_EXPORT_FAILED") if status.is_file() else {}
        code = str(receipt.get("error_code") or "WORD_EXPORT_FAILED")
        raise DeliveryError(code if re.fullmatch(r"[A-Z0-9_]+", code) else "WORD_EXPORT_FAILED")
    require(ascii_output.is_file() and ascii_output.stat().st_size > 0, "WORD_EXPORT_PDF_MISSING")
    receipt = load_json(status, "WORD_EXPORT_RECEIPT_INVALID")
    require(receipt.get("status") == "PASS_CUSTOMER_QA_WORD_PDF_EXPORT", "WORD_EXPORT_RECEIPT_NOT_PASS")
    write_bytes_atomic(pdf, ascii_output.read_bytes())
    require(sha256_file(pdf) == sha256_file(ascii_output), "WORD_EXPORT_PUBLISHED_PDF_HASH_MISMATCH")
    receipt["published"] = {
        "input_docx": artifact_ref(docx_path, output_dir).as_json(),
        "output_pdf": artifact_ref(pdf, output_dir).as_json(),
        "ascii_staging_required": True,
        "reason": "Windows PowerShell 5.1 child-process Unicode path compatibility",
    }
    write_json_atomic(status, receipt)
    return pdf, status


def build_checksums(output_dir: Path) -> Path:
    checksum_path = output_dir / "checksums.sha256"
    files = sorted(
        (path for path in output_dir.rglob("*") if path.is_file() and path != checksum_path),
        key=lambda path: relative_posix(path, output_dir).casefold(),
    )
    lines = [f"{sha256_file(path)}  {relative_posix(path, output_dir)}" for path in files]
    write_bytes_atomic(checksum_path, ("\n".join(lines) + "\n").encode("utf-8"))
    return checksum_path


def write_delivery_verification(output_dir: Path, scope: dict[str, Any], generated_at: str) -> Path:
    path = output_dir / "delivery_verification.json"
    artifacts = []
    for candidate in sorted((p for p in output_dir.rglob("*") if p.is_file()), key=lambda p: relative_posix(p, output_dir).casefold()):
        if candidate.name in {"delivery_verification.json", "checksums.sha256"}:
            continue
        artifacts.append(artifact_ref(candidate, output_dir).as_json())
    required = {
        "combined_scope.v1.json",
        "customer_research_evidence.xlsx",
        "客户技术研究报告.docx",
        "workbook_verification.json",
    }
    found = {item["path"] for item in artifacts}
    require(required.issubset(found), "DELIVERY_REQUIRED_OUTPUT_MISSING")
    value = {
        "schema": DELIVERY_SCHEMA,
        "status": "PASS",
        "generated_at_utc": generated_at,
        "scope_status": scope["status"],
        "company_id": scope["company_id"],
        "artifacts": artifacts,
        "checks": {
            "combined_scope_bound": True,
            "workbook_structural_and_preview_verification": "PASS",
            "docx_structural_verification": "PASS",
            "pdf_export": "PASS" if (output_dir / "客户技术研究报告.pdf").is_file() else "NOT_REQUESTED",
            "customer_body_read": False,
            "relationship_node_values_exported": False,
            "full_text_fts_opened": False,
            "dangerous_excel_prefixes_escaped": True,
        },
        "checksum_policy": "checksums.sha256 covers every delivery file except itself and is written after this verification receipt",
    }
    write_json_atomic(path, value)
    return path


def verify_delivery_output(output_dir: Path) -> dict[str, Any]:
    output_dir = output_dir.resolve()
    verification = load_json(output_dir / "delivery_verification.json", "DELIVERY_VERIFICATION_INVALID")
    require(verification.get("schema") == DELIVERY_SCHEMA and verification.get("status") == "PASS", "DELIVERY_VERIFICATION_NOT_PASS")
    checksum_path = output_dir / "checksums.sha256"
    require(checksum_path.is_file(), "DELIVERY_CHECKSUMS_MISSING")
    checked = 0
    try:
        for raw_line in checksum_path.read_text(encoding="utf-8").splitlines():
            if not raw_line.strip():
                continue
            match = re.fullmatch(r"([A-F0-9]{64})  (.+)", raw_line)
            require(match is not None, "DELIVERY_CHECKSUM_LINE_INVALID")
            expected, rel = match.groups()
            candidate = resolve_within(output_dir, validate_relative_artifact_path(rel, "DELIVERY_CHECKSUM_PATH_INVALID"), "DELIVERY_CHECKSUM_PATH_ESCAPE")
            require(candidate.is_file(), "DELIVERY_CHECKSUM_FILE_MISSING")
            require(sha256_file(candidate) == expected, "DELIVERY_CHECKSUM_MISMATCH")
            checked += 1
    except (OSError, UnicodeError) as exc:
        raise DeliveryError("DELIVERY_CHECKSUMS_INVALID") from exc
    require(checked > 0, "DELIVERY_CHECKSUMS_EMPTY")
    return {"status": "PASS", "checked_files": checked, "output_dir": str(output_dir)}


def choose_output_dir(case_root: Path, explicit: Path | None) -> Path:
    if explicit is not None:
        output = explicit.resolve()
    else:
        stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        output = (case_root / "deliveries" / "single_customer_final_delivery" / stamp).resolve()
    if output.exists():
        require(output.is_dir() and not any(output.iterdir()), "OUTPUT_DIRECTORY_NOT_EMPTY")
    else:
        output.mkdir(parents=True)
    return output


def build_delivery(args: argparse.Namespace) -> dict[str, Any]:
    case_root = args.case_root.resolve()
    output_dir = choose_output_dir(case_root, args.output_dir)
    generated_at = args.generated_at or now_utc()
    bundle = build_source_bundle(
        case_root,
        args.company_id,
        args.processing_scope,
        args.trade_receipt,
        args.trade_manifest,
        args.verify_manual_artifact_hashes,
    )
    scope = combined_scope(bundle, generated_at)
    scope_path = output_dir / "combined_scope.v1.json"
    write_json_atomic(scope_path, scope)

    workbook_path = output_dir / "customer_research_evidence.xlsx"
    workbook_verification = output_dir / "workbook_verification.json"
    build_workbook(
        workbook_payload(bundle, scope),
        workbook_path,
        output_dir / "qa" / "workbook",
        workbook_verification,
        args.node_exe,
        args.node_modules,
    )

    docx_path = output_dir / "客户技术研究报告.docx"
    build_docx(bundle, scope, docx_path)
    if args.export_pdf:
        export_pdf_with_word(docx_path, output_dir, args.company_id, args.word_timeout_seconds)

    write_delivery_verification(output_dir, scope, generated_at)
    build_checksums(output_dir)
    verified = verify_delivery_output(output_dir)
    return {
        "status": "PASS",
        "scope_status": scope["status"],
        "output_dir": str(output_dir),
        "checked_files": verified["checked_files"],
        "coverage_rows": len(bundle.coverage_rows),
        "relation_rows": len(bundle.relation_rows),
        "evidence_rows": len(bundle.evidence_rows),
        "pdf_exported": args.export_pdf,
    }


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description="Build a privacy-bounded OKKI single-customer final delivery package")
    sub = root.add_subparsers(dest="command", required=True)
    build = sub.add_parser("build", help="validate source receipts and build combined scope, XLSX, DOCX, optional PDF, hashes and verification")
    build.add_argument("--case-root", type=Path, required=True)
    build.add_argument("--company-id", required=True)
    build.add_argument("--processing-scope", type=Path)
    build.add_argument("--trade-receipt", type=Path)
    build.add_argument("--trade-manifest", type=Path)
    build.add_argument("--output-dir", type=Path)
    build.add_argument("--generated-at", help="fixed UTC timestamp for reproducible fixtures")
    build.add_argument("--verify-manual-artifact-hashes", action="store_true", help="rehash every manual-trade artifact; default validates manifest/receipt hash plus artifact paths and sizes")
    build.add_argument("--export-pdf", action="store_true", help="call the existing isolated Word COM exporter")
    build.add_argument("--word-timeout-seconds", type=int, default=600)
    build.add_argument("--node-exe", type=Path)
    build.add_argument("--node-modules", type=Path)

    verify = sub.add_parser("verify", help="recompute and verify delivery checksums without reading customer bodies")
    verify.add_argument("--output-dir", type=Path, required=True)
    return root


# EN: Parse CLI inputs and emit status-only results.
# 中文：解析命令行并仅输出状态结果。
def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        if args.command == "build":
            result = build_delivery(args)
        else:
            result = verify_delivery_output(args.output_dir)
        print(json.dumps(result, ensure_ascii=False, sort_keys=True))
        return 0
    except DeliveryError as exc:
        code = str(exc)
        if not re.fullmatch(r"[A-Z0-9_]+", code):
            code = "FINAL_DELIVERY_FAILED"
        print(code, file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
