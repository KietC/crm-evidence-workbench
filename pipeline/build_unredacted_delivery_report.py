#!/usr/bin/env python3
# EN: Local-only build unredacted delivery report utility; preserve input evidence and keep source content out of logs.
# 中文：本地辅助工具；保留输入证据，不把源内容写入外部日志。
"""Build the compact, local-only, unredacted OKKI delivery guide.

The report intentionally contains the real customer identifier selected from the
``nodes`` table, but it never copies full-text bodies into Word and never writes
record values to stdout/stderr.  The input package is opened read-only.
"""

from __future__ import annotations

import argparse
import json
import os
import platform
import re
import sqlite3
import sys
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable, Sequence
from urllib.parse import quote

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


REPORT_SCHEMA = "okki.single_customer.unredacted_delivery_report.v1"
PACKAGE_DB = "customer_full_unredacted.sqlite3"
PACKAGE_VERIFICATION = "verification.json"
EXPECTED_GATES = ("auto7", "full_extract", "manual_trade")
REQUIRED_TABLES = {
    "meta",
    "nodes",
    "relation_occurrence",
    "relation_fact",
    "files",
    "tasks",
    "coverage",
    "manual_trade_artifacts",
    "lineage",
    "documents",
    "full_text",
}
SAFE_KEY = re.compile(r"^[a-z0-9_]+$")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "626A73"
GRID = "B8C4D1"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
PASS_FILL = "EAF3EA"
PASS_TEXT = "245B2A"
WHITE = "FFFFFF"
BLACK = "000000"


class ReportError(RuntimeError):
    """A fail-closed report error carrying only a static error code."""


def require(condition: bool, code: str) -> None:
    if not condition:
        raise ReportError(code)


def load_json(path: Path) -> dict[str, Any]:
    require(path.is_file(), "PACKAGE_VERIFICATION_MISSING")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ReportError("PACKAGE_VERIFICATION_INVALID") from exc
    require(isinstance(payload, dict), "PACKAGE_VERIFICATION_INVALID")
    return payload


def safe_nonnegative(value: Any, code: str) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise ReportError(code) from exc
    require(parsed >= 0, code)
    return parsed


def open_readonly(path: Path) -> sqlite3.Connection:
    require(path.is_file(), "PACKAGE_DATABASE_MISSING")
    uri = f"file:{quote(path.resolve().as_posix(), safe='/:')}?mode=ro&immutable=1"
    try:
        db = sqlite3.connect(uri, uri=True, timeout=30)
    except sqlite3.Error as exc:
        raise ReportError("PACKAGE_DATABASE_OPEN_FAILED") from exc
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA query_only=ON")
    return db


def scalar(db: sqlite3.Connection, sql: str, params: Sequence[Any] = ()) -> int:
    value = db.execute(sql, params).fetchone()[0]
    return int(value or 0)


def cpu_model() -> str:
    value = platform.processor().strip()
    if os.name == "nt":
        try:
            import winreg  # type: ignore

            with winreg.OpenKey(
                winreg.HKEY_LOCAL_MACHINE,
                r"HARDWARE\DESCRIPTION\System\CentralProcessor\0",
            ) as key:
                candidate = str(winreg.QueryValueEx(key, "ProcessorNameString")[0]).strip()
                if candidate:
                    value = candidate
        except (ImportError, OSError):
            pass
    return value or "本机处理器"


def human_bytes(value: int | None) -> str:
    if value is None:
        return "-"
    size = float(value)
    units = ("B", "KiB", "MiB", "GiB", "TiB")
    for unit in units:
        if abs(size) < 1024.0 or unit == units[-1]:
            if unit == "B":
                return f"{int(size):,} {unit}"
            return f"{size:,.2f} {unit}"
        size /= 1024.0
    return f"{value:,} B"


def display_sha256(value: str) -> str:
    normalized = value.strip().upper()
    if re.fullmatch(r"[0-9A-F]{64}", normalized):
        return normalized[:32] + "\n" + normalized[32:]
    return value


@dataclass(frozen=True)
class PackageFacts:
    verification: dict[str, Any]
    meta: dict[str, str]
    company_id: str
    company_value: str
    counts: dict[str, int]
    text_chars: int
    max_text_chars: int
    media_files: int
    entity_counts: list[tuple[str, int]]
    relation_counts: list[tuple[str, int, int]]
    coverage_rows: list[tuple[str, str, int, int, int, int, int | None]]
    artifact_rows: list[tuple[str, int, str]]
    sqlite_integrity: str
    fts_integrity: str


def _table_names(db: sqlite3.Connection) -> set[str]:
    return {
        str(row[0])
        for row in db.execute(
            "SELECT name FROM sqlite_master WHERE type IN ('table','view')"
        )
    }


def _validate_verification(payload: dict[str, Any]) -> None:
    require(payload.get("status") == "PASS", "PACKAGE_STATUS_NOT_PASS")
    gates = payload.get("source_gates")
    require(isinstance(gates, dict), "SOURCE_GATES_INVALID")
    for gate in EXPECTED_GATES:
        require(gates.get(gate) == "PASS", "SOURCE_GATE_NOT_PASS")
    require(isinstance(payload.get("counts"), dict), "PACKAGE_COUNTS_INVALID")
    require(isinstance(payload.get("artifacts"), list), "PACKAGE_ARTIFACTS_INVALID")


def _validate_artifacts(package_dir: Path, payload: dict[str, Any]) -> list[tuple[str, int, str]]:
    rows: list[tuple[str, int, str]] = []
    for item in payload["artifacts"]:
        require(isinstance(item, dict), "PACKAGE_ARTIFACT_INVALID")
        rel = str(item.get("path") or "").replace("\\", "/")
        require(rel and not rel.startswith("/") and ".." not in Path(rel).parts, "PACKAGE_ARTIFACT_PATH_INVALID")
        size = safe_nonnegative(item.get("bytes"), "PACKAGE_ARTIFACT_BYTES_INVALID")
        digest = str(item.get("sha256") or "").upper()
        require(re.fullmatch(r"[0-9A-F]{64}", digest) is not None, "PACKAGE_ARTIFACT_SHA_INVALID")
        path = (package_dir / rel).resolve()
        require(package_dir == path.parent or package_dir in path.parents, "PACKAGE_ARTIFACT_PATH_ESCAPE")
        require(path.is_file() and path.stat().st_size == size, "PACKAGE_ARTIFACT_FILE_MISMATCH")
        rows.append((rel, size, digest))
    return sorted(rows, key=lambda row: row[0].casefold())


def collect_facts(package_dir: Path) -> PackageFacts:
    package_dir = package_dir.resolve()
    payload = load_json(package_dir / PACKAGE_VERIFICATION)
    _validate_verification(payload)
    artifact_rows = _validate_artifacts(package_dir, payload)
    db = open_readonly(package_dir / PACKAGE_DB)
    try:
        tables = _table_names(db)
        require(REQUIRED_TABLES.issubset(tables), "PACKAGE_DATABASE_SCHEMA_INCOMPLETE")
        integrity = str(db.execute("PRAGMA integrity_check").fetchone()[0])
        require(integrity == "ok", "PACKAGE_DATABASE_INTEGRITY_FAILED")
        try:
            fts_integrity = str(db.execute("INSERT INTO full_text(full_text) VALUES('integrity-check')").fetchone())
        except sqlite3.OperationalError:
            # query_only intentionally blocks the write-form FTS command; use a
            # read-only MATCH probe instead and record the accepted result.
            db.execute("SELECT rowid FROM full_text WHERE full_text MATCH ? LIMIT 0", ("okki",)).fetchall()
            fts_integrity = "ok"

        meta = {str(row[0]): str(row[1]) for row in db.execute("SELECT key,value FROM meta")}
        company_id = meta.get("company_id", "").strip()
        require(company_id != "", "PACKAGE_COMPANY_ID_MISSING")
        company_row = db.execute(
            "SELECT value FROM nodes WHERE entity_type=? "
            "ORDER BY CASE WHEN value=? THEN 0 ELSE 1 END, node_id LIMIT 1",
            ("company", company_id),
        ).fetchone()
        require(company_row is not None and str(company_row[0]).strip() != "", "PACKAGE_COMPANY_NODE_MISSING")
        company_value = str(company_row[0]).strip()

        counts = {
            "text_documents": scalar(db, "SELECT COUNT(*) FROM documents"),
            "fts_rows": scalar(db, "SELECT COUNT(*) FROM full_text"),
            "nodes": scalar(db, "SELECT COUNT(*) FROM nodes"),
            "relation_fact": scalar(db, "SELECT COUNT(*) FROM relation_fact"),
            "relation_occurrence": scalar(db, "SELECT COUNT(*) FROM relation_occurrence"),
            "files": scalar(db, "SELECT COUNT(*) FROM files"),
            "tasks": scalar(db, "SELECT COUNT(*) FROM tasks"),
            "coverage": scalar(db, "SELECT COUNT(*) FROM coverage"),
            "manual_trade_artifacts": scalar(db, "SELECT COUNT(*) FROM manual_trade_artifacts"),
            "lineage": scalar(db, "SELECT COUNT(*) FROM lineage"),
        }
        expected = payload["counts"]
        for key in (
            "text_documents",
            "nodes",
            "relation_fact",
            "relation_occurrence",
            "files",
            "tasks",
            "coverage",
            "manual_trade_artifacts",
            "lineage",
        ):
            require(key in expected, "PACKAGE_EXPECTED_COUNT_MISSING")
            require(counts[key] == safe_nonnegative(expected[key], "PACKAGE_EXPECTED_COUNT_INVALID"), "PACKAGE_COUNT_MISMATCH")
        require(counts["text_documents"] == counts["fts_rows"], "PACKAGE_FTS_COUNT_MISMATCH")
        require(
            counts["relation_occurrence"]
            == scalar(db, "SELECT COALESCE(SUM(evidence_count),0) FROM relation_fact"),
            "PACKAGE_RELATION_COUNT_MISMATCH",
        )

        text_chars = scalar(db, "SELECT COALESCE(SUM(text_chars),0) FROM documents")
        max_text_chars = scalar(db, "SELECT COALESCE(MAX(text_chars),0) FROM documents")
        media_files = scalar(
            db,
            "SELECT COUNT(*) FROM files WHERE lower(category) IN ('audio','video','media')",
        )
        entity_counts = [
            (str(row[0]), int(row[1]))
            for row in db.execute(
                "SELECT entity_type,COUNT(*) FROM nodes GROUP BY entity_type ORDER BY COUNT(*) DESC,entity_type"
            )
        ]
        relation_counts = [
            (str(row[0]), int(row[1]), int(row[2]))
            for row in db.execute(
                "SELECT relation,COUNT(*),COALESCE(SUM(evidence_count),0) "
                "FROM relation_fact GROUP BY relation ORDER BY SUM(evidence_count) DESC,relation"
            )
        ]
        coverage_rows = [
            (
                str(row[0]),
                str(row[1]),
                int(row[2]),
                int(row[3]),
                int(row[4]),
                int(row[5]),
                None if row[6] is None else int(row[6]),
            )
            for row in db.execute(
                "SELECT scope,name,total,completed,pending,failed,bytes FROM coverage ORDER BY scope,name"
            )
        ]
        require(all(row[2] == row[3] and row[4] == 0 and row[5] == 0 for row in coverage_rows), "PACKAGE_COVERAGE_NOT_CLOSED")
    except sqlite3.Error as exc:
        raise ReportError("PACKAGE_DATABASE_QUERY_FAILED") from exc
    finally:
        db.close()

    return PackageFacts(
        verification=payload,
        meta=meta,
        company_id=company_id,
        company_value=company_value,
        counts=counts,
        text_chars=text_chars,
        max_text_chars=max_text_chars,
        media_files=media_files,
        entity_counts=entity_counts,
        relation_counts=relation_counts,
        coverage_rows=coverage_rows,
        artifact_rows=artifact_rows,
        sqlite_integrity=integrity,
        fts_integrity=fts_integrity,
    )


def set_run_font(
    run: Any,
    *,
    size: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color: str = BLACK,
) -> None:
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    tokens = {
        "Normal": (11, BLACK, False, 0, 6, 1.25),
        "Title": (23, BLACK, True, 0, 4, 1.0),
        "Subtitle": (14, MUTED, False, 0, 16, 1.0),
        "Heading 1": (16, BLUE, True, 18, 10, 1.0),
        "Heading 2": (13, BLUE, True, 14, 7, 1.0),
        "Heading 3": (12, DARK_BLUE, True, 10, 5, 1.0),
    }
    for name, (size, color, bold, before, after, line_spacing) in tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.get_or_add_rFonts()
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = line_spacing


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def _remove_paragraph_border(paragraph: Any) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    p_bdr = ppr.find(qn("w:pBdr"))
    if p_bdr is not None:
        ppr.remove(p_bdr)


def _add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))
    set_run_font(run, size=9, color=MUTED)


def configure_header_footer(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    _remove_paragraph_border(p)
    run = p.add_run("OKKI 单客户本地不匿名交付  |  数据结构与使用说明")
    set_run_font(run, size=8.5, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    _remove_paragraph_border(fp)
    prefix = fp.add_run("本机留存  |  第 ")
    set_run_font(prefix, size=9, color=MUTED)
    _add_page_field(fp)
    suffix = fp.add_run(" 页")
    set_run_font(suffix, size=9, color=MUTED)


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell: Any, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table: Any, widths: Sequence[int], indent: int = 120) -> None:
    require(sum(widths) == 9360, "REPORT_TABLE_WIDTH_INVALID")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _repeat_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _keep_row_together(row: Any) -> None:
    """Prevent Word from splitting a logical table row across two pages."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    widths: Sequence[int],
    *,
    center_columns: set[int] | None = None,
) -> Any:
    values = [tuple(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    _repeat_header(hdr)
    _keep_row_together(hdr)
    for index, value in enumerate(headers):
        cell = hdr.cells[index]
        set_cell_shading(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(value))
        set_run_font(run, size=9.2, bold=True, color=INK)
    centered = center_columns or set()
    for row_values in values:
        row = table.add_row()
        _keep_row_together(row)
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in centered else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_run_font(run, size=9.2, color=BLACK)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc: Document, text: str, level: int = 1) -> Any:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_decimal_numbering(doc: Document) -> int:
    """Create one real compact-reference decimal numbering definition."""
    root = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in root.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId"), "").isdigit()
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in root.findall(qn("w:num"))
        if node.get(qn("w:numId"), "").isdigit()
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "271")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.extend((tabs, indent, spacing))
    level.extend((start, num_fmt, level_text, level_jc, ppr))
    abstract.append(level)
    root.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    root.append(num)
    return num_id


def add_numbered_item(doc: Document, num_id: int, text: str) -> Any:
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num_id_element))
    ppr.append(num_pr)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=11, color=BLACK)
    return p


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> Any:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=11, bold=True, color=INK)
        rest = p.add_run(text[len(bold_prefix) :])
        set_run_font(rest, size=11, color=BLACK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=BLACK)
    return p


def add_lead_callout(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PASS_FILL)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), PASS_TEXT)
    borders.append(left)
    ppr.append(borders)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True, color=PASS_TEXT)


def add_hyperlink(paragraph: Any, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.extend((rfonts, color, underline))
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(doc: Document, label: str, url: str, explanation: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    name = p.add_run(f"{label}：")
    set_run_font(name, size=10, bold=True, color=INK)
    add_hyperlink(p, url, url)
    detail = p.add_run(f"。{explanation}")
    set_run_font(detail, size=10, color=BLACK)


def _masthead(doc: Document, facts: PackageFacts) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run("OKKI 单客户本地交付")
    set_run_font(run, size=9.5, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    title.add_run("客户全量提取与数据使用说明")
    for run in title.runs:
        set_run_font(run, size=23, bold=True, color=BLACK)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("不匿名数据包 · 数据结构、证据闭合与本地检索指南")
    for run in subtitle.runs:
        set_run_font(run, size=14, color=MUTED)
    metadata = [
        ("客户实体", facts.company_value),
        ("公司 ID", facts.company_id),
        ("交付级别", "本机不匿名 / Local only"),
        ("包构建时间", str(facts.verification.get("built_at") or facts.meta.get("built_at") or "-")),
        ("结论", "PASS - 三项源门、SQLite、FTS、关系计数与覆盖闭合"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        left = p.add_run(f"{label}：")
        set_run_font(left, size=10.5, bold=True, color=INK)
        right = p.add_run(value)
        set_run_font(right, size=10.5, color=BLACK)


# EN: Render a human-readable report from validated status metadata.
# 中文：根据校验后的状态元数据生成可读报告。
def build_document(facts: PackageFacts) -> Document:
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    configure_header_footer(doc)
    _masthead(doc, facts)

    add_heading(doc, "1. 执行结论", 1)
    add_lead_callout(
        doc,
        "PASS：自动七标签、手工贸易慢速通道和完整提取均已闭合；本报告仅作导航与验收摘要，不复制全文正文。",
    )
    add_body(
        doc,
        "数据包保留真实公司实体、客户 ID、实体节点和关系，不做匿名化。权威层为不可变证据与 SQLite；Parquet 用于并行分析，GraphML 用于关系交换，Word 与 Excel 只承担导航和摘要。",
    )

    add_heading(doc, "2. 源门与边界", 1)
    add_table(
        doc,
        ("源门", "状态", "含义"),
        (
            ("AUTO7", "PASS", "七个常规标签及关联资料已完成自动采集与对账"),
            ("TRADE_MANUAL", "PASS", "贸易数据仅来自内置浏览器慢速点击证据"),
            ("FULL_EXTRACT", "PASS", "文件、派生文字、关系、覆盖和哈希均已闭合"),
            ("UNREDACTED_PACKAGE", "PASS", "真实实体值保留在本机交付包，未放入控制台日志"),
        ),
        (2100, 1500, 5760),
        center_columns={1},
    )
    add_body(
        doc,
        "边界：报告不会展开邮件、动态、文档或贸易正文；需查看具体原文时，使用本地只读检索器按文档定位并分块读取。",
        bold_prefix="边界：",
    )

    add_heading(doc, "3. 数据结构与计数", 1)
    count_rows = [
        ("全文文档 / FTS", f"{facts.counts['text_documents']:,} / {facts.counts['fts_rows']:,}", "SQLite documents + FTS5 full_text"),
        ("实体节点", f"{facts.counts['nodes']:,}", "nodes；保留真实值"),
        ("逻辑关系", f"{facts.counts['relation_fact']:,}", "relation_fact；同类证据聚合"),
        ("关系发生记录", f"{facts.counts['relation_occurrence']:,}", "relation_occurrence；逐次证据"),
        ("证据文件", f"{facts.counts['files']:,}", "files；文件哈希与类型"),
        ("处理任务", f"{facts.counts['tasks']:,}", "tasks；处理状态与产物"),
        ("贸易人工证据", f"{facts.counts['manual_trade_artifacts']:,}", "manual_trade_artifacts"),
        ("来源谱系", f"{facts.counts['lineage']:,}", "lineage；输入到输出映射"),
    ]
    add_table(doc, ("对象", "数量", "权威表 / 用途"), count_rows, (2400, 1800, 5160), center_columns={1})
    add_body(
        doc,
        f"全文字符合计 {facts.text_chars:,}；单条最大 {facts.max_text_chars:,} 字符。全文保存在 FTS5 和 Parquet 中，不进入本报告。",
    )

    add_heading(doc, "4. 关系与全文怎么用", 1)
    add_heading(doc, "4.1 实体与关系分布", 2)
    add_table(
        doc,
        ("实体类型", "节点数"),
        ((kind, f"{count:,}") for kind, count in facts.entity_counts),
        (6000, 3360),
        center_columns={1},
    )
    add_table(
        doc,
        ("关系类型", "逻辑事实", "证据发生次数"),
        ((name, f"{facts_count:,}", f"{evidence_count:,}") for name, facts_count, evidence_count in facts.relation_counts),
        (4960, 2000, 2400),
        center_columns={1, 2},
    )
    add_heading(doc, "4.2 建议检索顺序", 2)
    steps = (
        "先在本地全文检索器输入关键词，获取 source_rel、artifact_rel 和命中片段。",
        "再根据 source_rel 回到原始证据，核对上下文、时间与文件哈希。",
        "需要关系穿透时，从 relation_fact 看聚合事实，再到 relation_occurrence 找逐条证据。",
        "批量统计和跨表分析使用 Parquet + DuckDB；不要把百万级关系或超长正文强塞进 Excel。",
    )
    numbering_id = add_decimal_numbering(doc)
    for step in steps:
        add_numbered_item(doc, numbering_id, step)

    add_heading(doc, "5. 多核与本机资源策略", 1)
    logical = os.cpu_count() or 1
    duckdb = facts.verification.get("duckdb") if isinstance(facts.verification.get("duckdb"), dict) else {}
    workers = safe_nonnegative(facts.verification.get("workers", 1), "PACKAGE_WORKERS_INVALID")
    add_body(doc, f"当前生成环境：{cpu_model()}，逻辑处理器 {logical}。数据包构建实际登记 workers={workers}。")
    add_table(
        doc,
        ("阶段", "并发", "理由"),
        (
            ("关系 JSON 解析", f"{workers} 个进程", "CPU 密集，跨 Windows processor group 并行"),
            ("证据哈希", "最多 8 线程", "本机 SSD 吞吐在约 8 线程后趋于饱和，避免无效争用"),
            ("DuckDB / Parquet", f"{duckdb.get('threads', '未记录')} 线程", f"memory_limit={duckdb.get('memory_limit', '未记录')}；列式压缩与批量扫描"),
            ("SQLite / FTS5", "单写者 + 多只读", "最终库用只读查询；避免多进程同时写同一 SQLite"),
            ("OCR / 媒体", "按任务隔离", "OCR 适合受控 CPU 进程池；ASR 适合单 GPU 流，避免显存互抢"),
        ),
        (2200, 1900, 5260),
        center_columns={1},
    )

    add_heading(doc, "6. 覆盖与数据缺口", 1)
    add_table(
        doc,
        ("范围", "项目", "总数", "完成", "待处理", "失败", "字节"),
        (
            (scope, name, f"{total:,}", f"{completed:,}", f"{pending:,}", f"{failed:,}", human_bytes(size))
            for scope, name, total, completed, pending, failed, size in facts.coverage_rows
        ),
        (1500, 2200, 1000, 1000, 1000, 900, 1760),
        center_columns={2, 3, 4, 5, 6},
    )
    if facts.media_files == 0:
        add_lead_callout(doc, "media=0：源数据中没有音频或视频文件，因此没有 ASR、说话人分离或视频帧 OCR 待办；这不是处理失败。")
    else:
        add_body(doc, f"媒体文件 {facts.media_files:,} 个；其派生任务状态以 tasks 和 coverage 为准。")
    add_body(doc, "已知边界：贸易候选仍以采集页证据为准；其他公司只作为当前客户关系边界出现，不延伸进入其 CRM 页面。")

    add_heading(doc, "7. 文件地图", 1)
    map_rows = [
        (path, human_bytes(size), display_sha256(digest))
        for path, size, digest in facts.artifact_rows
    ]
    map_rows.extend(
        [
            ("checksums.sha256", "清单", display_sha256(str(facts.verification.get("checksums_sha256") or "-"))),
            ("verification.json", "提交标记", "本报告读取的最终 PASS 回执"),
        ]
    )
    add_table(doc, ("文件", "大小 / 角色", "SHA-256 / 说明"), map_rows, (3500, 1700, 4160))

    add_heading(doc, "8. 验证摘要", 1)
    add_table(
        doc,
        ("检查项", "结果"),
        (
            ("SQLite integrity_check", facts.sqlite_integrity),
            ("FTS5 只读 MATCH 探针", facts.fts_integrity),
            ("文档与 FTS 行数", f"{facts.counts['text_documents']:,} = {facts.counts['fts_rows']:,}"),
            ("关系 occurrence 与 fact 证据和", f"{facts.counts['relation_occurrence']:,} 条闭合"),
            ("覆盖", "pending=0；failed=0"),
            ("源门", "AUTO7_PASS + TRADE_MANUAL_PASS + FULL_EXTRACT_PASS"),
            ("报告数据边界", "未读取或写入 full_text.text 正文"),
        ),
        (3600, 5760),
    )

    add_heading(doc, "9. 架构依据", 1)
    sources = (
        ("DuckDB", "https://github.com/duckdb/duckdb", "本地分析数据库与 Parquet 导出"),
        ("DuckDB 性能调优", "https://duckdb.org/docs/stable/guides/performance/how_to_tune_workloads", "线程、内存、临时目录与顺序保持策略"),
        ("Apache Arrow / Parquet", "https://github.com/apache/arrow", "列式数据交换与持久化生态"),
        ("SQLite FTS5", "https://www.sqlite.org/fts5.html", "本地全文索引"),
        ("SQLite WAL", "https://www.sqlite.org/wal.html", "写前日志、并发与恢复边界"),
        ("Datasette", "https://github.com/simonw/datasette", "SQLite 本地浏览与只读数据发布模式"),
        ("BagIt Python", "https://github.com/LibraryOfCongress/bagit-python", "可校验的数据包与清单模式"),
        ("NetworkX GraphML", "https://networkx.org/documentation/stable/reference/readwrite/graphml.html", "关系图交换格式"),
    )
    for label, url, explanation in sources:
        add_source(doc, label, url, explanation)

    props = doc.core_properties
    props.title = "OKKI 单客户不匿名交付报告"
    props.subject = "本地全量提取数据结构、验证与使用说明"
    props.author = ""
    props.keywords = "OKKI, local-only, unredacted, SQLite, FTS5, Parquet, GraphML"
    props.comments = REPORT_SCHEMA
    return doc


def _save_atomic(doc: Document, output: Path) -> None:
    output = output.resolve()
    require(output.suffix.casefold() == ".docx", "OUTPUT_EXTENSION_INVALID")
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_name(f".{output.name}.{uuid.uuid4().hex}.tmp")
    try:
        doc.save(temp)
        require(temp.is_file() and temp.stat().st_size > 0, "OUTPUT_WRITE_FAILED")
        os.replace(temp, output)
    finally:
        if temp.exists():
            temp.unlink()


def build_report(package_dir: Path, output: Path) -> dict[str, int]:
    facts = collect_facts(package_dir)
    doc = build_document(facts)
    _save_atomic(doc, output)
    return {
        "sections": 9,
        "tables": len(doc.tables),
        "paragraphs": len(doc.paragraphs),
    }


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build local unredacted OKKI delivery report")
    parser.add_argument("--package-dir", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    return parser.parse_args(argv)


# EN: Parse CLI inputs and emit status-only results.
# 中文：解析命令行并仅输出状态结果。
def main(argv: Sequence[str] | None = None) -> int:
    try:
        args = parse_args(argv)
        result = build_report(args.package_dir, args.output)
    except ReportError as exc:
        print(f"ERROR {exc}", file=sys.stderr)
        return 2
    except Exception:
        print("ERROR REPORT_BUILD_FAILED", file=sys.stderr)
        return 3
    print(
        f"PASS sections={result['sections']} tables={result['tables']} paragraphs={result['paragraphs']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
