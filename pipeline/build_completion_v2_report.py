#!/usr/bin/env python3
# EN: Local-only build completion v2 report utility; preserve input evidence and keep source content out of logs.
# 中文：本地辅助工具；保留输入证据，不把源内容写入外部日志。
"""Build the compact local-only completion-v2 Chinese report.

The report uses the established unredacted delivery visual system.  It only
contains status, counts, gap classes and audit guidance; customer bodies stay
in the local SQLite/FTS package and are never written to stdout.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any, Sequence

from docx import Document
from docx.shared import Pt

from build_unredacted_delivery_report import (
    BLACK,
    BLUE,
    INK,
    MUTED,
    _save_atomic,
    add_body,
    add_decimal_numbering,
    add_heading,
    add_lead_callout,
    add_numbered_item,
    add_table,
    configure_header_footer,
    configure_page,
    configure_styles,
    set_run_font,
)


SCHEMA = "okki.single_customer.completion_v2_catalog_payload.v1"


class ReportError(RuntimeError):
    pass


def require(condition: bool, code: str) -> None:
    if not condition:
        raise ReportError(code)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(4 * 1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def read_payload(path: Path) -> dict[str, Any]:
    require(path.is_file(), "REPORT_PAYLOAD_MISSING")
    try:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ReportError("REPORT_PAYLOAD_INVALID") from exc
    require(isinstance(value, dict) and value.get("schema") == SCHEMA, "REPORT_PAYLOAD_SCHEMA_INVALID")
    require(isinstance(value.get("counts"), dict), "REPORT_COUNTS_INVALID")
    require(isinstance(value.get("gates"), dict), "REPORT_GATES_INVALID")
    require(isinstance(value.get("source_gap_rows"), list), "REPORT_SOURCE_GAPS_INVALID")
    require(isinstance(value.get("diff_rows"), list), "REPORT_DIFF_INVALID")
    require(value.get("accessible_data_status") == "ACCESSIBLE_DATA_PASS", "REPORT_ACCESSIBLE_STATUS_INVALID")
    require(value.get("absolute_completeness_status") in {"ABSOLUTE_COMPLETENESS_SOURCE_GAPS", "ABSOLUTE_COMPLETENESS_PASS"}, "REPORT_ABSOLUTE_STATUS_INVALID")
    return value


def human_int(value: Any) -> str:
    try:
        return f"{int(value):,}"
    except (TypeError, ValueError) as exc:
        raise ReportError("REPORT_COUNT_INVALID") from exc


# EN: Render a human-readable report from validated status metadata.
# 中文：根据校验后的状态元数据生成可读报告。
def build_document(payload: dict[str, Any]) -> Document:
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    configure_header_footer(doc)
    doc.core_properties.title = f"客户{payload['company_id']}全案例递归补全研究报告v2"
    doc.core_properties.subject = "OKKI 单客户本地不匿名递归补全"

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(2)
    set_run_font(kicker.add_run("OKKI 单客户本地交付"), size=9.5, bold=True, color=BLUE)
    title = doc.add_paragraph(style="Title")
    set_run_font(title.add_run("全案例递归补全研究报告 v2"), size=23, bold=True, color=BLACK)
    subtitle = doc.add_paragraph(style="Subtitle")
    set_run_font(subtitle.add_run("原件递归 · 显式关系 · 双时间线 · 客观缺口"), size=14, color=MUTED)
    for label, value in (
        ("公司 ID", payload["company_id"]),
        ("交付级别", "本机不匿名 / Local only"),
        ("生成时间", str(payload.get("generated_at_utc") or "-")),
        ("可取得数据", payload["accessible_data_status"]),
        ("绝对完整", payload["absolute_completeness_status"]),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_run_font(paragraph.add_run(f"{label}："), size=10.5, bold=True, color=INK)
        set_run_font(paragraph.add_run(str(value)), size=10.5, color=BLACK)

    add_heading(doc, "1. 执行结论", 1)
    add_lead_callout(doc, "ACCESSIBLE_DATA_PASS：所有当前可取得内容已完成处理、关联和验收；旧版已验收成果未覆盖、未重算。")
    if payload["absolute_completeness_status"] == "ABSOLUTE_COMPLETENESS_SOURCE_GAPS":
        add_body(doc, "ABSOLUTE_COMPLETENESS_SOURCE_GAPS：源端删除、失效URL、缺少密码或未返回内容继续作为客观缺口保留，不能写成绝对完整。", bold_prefix="ABSOLUTE_COMPLETENESS_SOURCE_GAPS：")
    else:
        add_body(doc, "ABSOLUTE_COMPLETENESS_PASS：本轮未登记客观源缺口。", bold_prefix="ABSOLUTE_COMPLETENESS_PASS：")
    add_body(doc, "Word与Excel仅承担结论、目录和审计导航。完整正文、OCR、ASR、关系证据与时间事件以本地SQLite、FTS、Parquet、GraphML和原件库为权威。")

    counts = payload["counts"]
    add_heading(doc, "2. 补全规模", 1)
    add_table(
        doc,
        ("指标", "数量", "验收意义"),
        (
            ("旧版合格文件", human_int(counts["old_files"]), "原成果保留"),
            ("v2证据文件", human_int(counts["v2_files"]), "新增附件与补采物已入范围"),
            ("历史邮件附件原件", human_int(counts["historical_mail_files"]), "字节与SHA绑定"),
            ("内容对象", human_int(counts["content_objects"]), "按SHA去重"),
            ("出现实例", human_int(counts["occurrences"]), "物理位置与父级不丢失"),
            ("处理任务", human_int(counts["tasks"]), "OCR/解析/ASR终态"),
            ("FTS文档", human_int(counts["fts_documents"]), "本地全文检索"),
        ),
        (3000, 1600, 4760),
        center_columns={1},
    )

    add_heading(doc, "3. 关系与时间顺序", 1)
    add_table(
        doc,
        ("关系域", "数量", "口径"),
        (
            ("邮件附件出现记录", human_int(counts["attachment_occurrences"]), "每次实际引用保留"),
            ("唯一邮件附件关系", human_int(counts["mail_attachment_relations"]), "邮件→附件去重事实"),
            ("多父附件", human_int(counts["multiparent_attachments"]), "同一内容对应多封邮件"),
            ("显式关系", human_int(counts["explicit_relations"]), "仅源字段支持的结构/业务关系"),
            ("业务时间事件", human_int(counts["business_timeline"]), "发生时间；无时区不猜UTC"),
            ("采集时间事件", human_int(counts["capture_timeline"]), "证据采集与处理顺序"),
        ),
        (3000, 1600, 4760),
        center_columns={1},
    )
    add_body(doc, "旧版first_source_rel/last_source_rel路径字典序不再充当时间顺序。v2保留原始时间文本、解析状态、精度、时区信息、来源路径、JSON指针和稳定排序键。")

    add_heading(doc, "4. 验收门槛", 1)
    numbering = add_decimal_numbering(doc)
    labels = {
        "PAGE_GAP_RECOVERY_PASS": "页面缺口补采闭合：其他动态及实际可见筛选状态有独立证据。",
        "MAIL_ATTACHMENT_SCOPE_PASS": "邮件附件范围闭合：历史原件、出现记录和多父引用数量守恒。",
        "CONTENT_RECURSION_PASS_WITH_SOURCE_GAPS": "内容递归完成：可取得对象终态化，客观不可取得内容进入source_gaps。",
        "RELATIONSHIP_V2_PASS": "关系v2通过：显式边零孤儿、证据可回溯。",
        "TIMELINE_V2_PASS": "双时间线通过：原始时间零丢失，业务与采集顺序可重建。",
        "DELIVERY_V2_PASS": "交付验收在Excel、DOCX、PDF及全部渲染检查通过后提交。",
    }
    for gate, text in labels.items():
        state = payload["gates"].get(gate) is True
        add_numbered_item(doc, numbering, f"{'PASS' if state else 'FINALIZE时提交'} · {text}")

    add_heading(doc, "5. 客观源缺口", 1)
    gap_rows = payload["source_gap_rows"]
    gap_counter = Counter(str(row[3] or "UNSPECIFIED") for row in gap_rows if isinstance(row, list) and len(row) >= 4)
    if gap_rows:
        add_table(
            doc,
            ("错误码", "数量", "处理口径"),
            ((code, human_int(count), "保留原始指针和哈希；不猜写、不伪装成功") for code, count in sorted(gap_counter.items(), key=lambda item: (-item[1], item[0]))),
            (3600, 1400, 4360),
            center_columns={1},
        )
        add_body(doc, f"本轮共登记{len(gap_rows):,}项客观源缺口。完整逐项清单位于交付目录source_gaps.csv、source_gaps.json和权威SQLite。")
    else:
        add_body(doc, "本轮未登记客观源缺口。")

    add_heading(doc, "6. 新旧差异", 1)
    add_table(
        doc,
        ("指标", "旧版", "补全版", "增量"),
        ((str(row[0]), human_int(row[1]), human_int(row[2]), human_int(row[3])) for row in payload["diff_rows"]),
        (3960, 1800, 1800, 1800),
        center_columns={1, 2, 3},
    )

    add_heading(doc, "7. 本地使用与权威顺序", 1)
    for index, item in enumerate((
        "先核对verification.json与checksums.sha256；只有ACCESSIBLE_DATA_PASS才视为交付闭合。",
        "原件以evidence/original为最高权威；同内容SHA去重不删除出现实例和父级引用。",
        "全文检索使用full_extract_v2/full_text.sqlite3；不要把客户正文上传外部模型。",
        "关系查询使用relationship_v2/relationship_timeline_v2.sqlite3或Parquet；GraphML用于图工具交换。",
        "Excel若显示截断，只代表工作簿展示上限；完整计数和明细仍在SQLite/Parquet。",
    ), start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(14)
        paragraph.paragraph_format.first_line_indent = Pt(-14)
        paragraph.paragraph_format.space_after = Pt(5)
        set_run_font(paragraph.add_run(f"{index}.  {item}"), size=10.5, color=BLACK)

    add_heading(doc, "8. 结论边界", 1)
    add_body(doc, "ACCESSIBLE_DATA_PASS不等于源世界绝对完整。当前可访问内容已经闭合；任何源端删除、失效、权限不足、需密码或未返回内容均保留在ABSOLUTE_COMPLETENESS_SOURCE_GAPS中。")
    return doc


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description="Build completion-v2 local Chinese DOCX report")
    result.add_argument("--payload", type=Path, required=True)
    result.add_argument("--output", type=Path)
    result.add_argument("--validate-only", action="store_true")
    return result


# EN: Parse CLI inputs and emit status-only results.
# 中文：解析命令行并仅输出状态结果。
def main(argv: Sequence[str] | None = None) -> int:
    args = parser().parse_args(argv)
    try:
        payload = read_payload(args.payload)
        if args.validate_only:
            print(json.dumps({"status": "VALID", "schema": payload["schema"], "source_gaps": len(payload["source_gap_rows"]), "diff_rows": len(payload["diff_rows"])}, separators=(",", ":")))
            return 0
        require(args.output is not None, "REPORT_OUTPUT_MISSING")
        document = build_document(payload)
        _save_atomic(document, args.output.resolve())
        print(json.dumps({"status": "PASS", "bytes": args.output.stat().st_size, "sha256": sha256_file(args.output)}, separators=(",", ":")))
        return 0
    except ReportError as exc:
        print(json.dumps({"status": "FAILED", "error_code": str(exc)}, separators=(",", ":")), file=sys.stderr)
        return 2
    except Exception:
        print(json.dumps({"status": "FAILED", "error_code": "REPORT_BUILD_FAILED"}, separators=(",", ":")), file=sys.stderr)
        return 3


if __name__ == "__main__":
    raise SystemExit(main())
